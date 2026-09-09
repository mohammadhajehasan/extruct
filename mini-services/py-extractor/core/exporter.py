# -*- coding: utf-8 -*-
"""
core/exporter.py — المصدّر: CSV utf-8-sig + Excel RTL منسق (openpyxl).
المصدر: الخطة الحاكمة v7.1 — الجزء 9 + عقد API (/api/export).
XLSX: RTL، ترويسة fill 1F4E78 بخط أبيض عريض، حدود، عرض تلقائي، تجميد A2.
"""
import csv
import io
import time
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MECHANIC_EXTRA_HEADERS = ["الفئة", "المصدر"]


def export_csv(headers: List[str], rows: List[List[str]]) -> bytes:
    """CSV بترميز utf-8-sig (متوافق مع Excel العربي)."""
    buf = io.StringIO()
    w = csv.writer(buf)  # فواصل قياسية
    if headers:
        w.writerow(headers)
    for r in (rows or []):
        w.writerow(["" if c is None else str(c) for c in r])
    return buf.getvalue().encode("utf-8-sig")


def _auto_width(ws, headers: List[str], rows: List[List[str]]) -> None:
    """عرض أعمدة تلقائي (أطول قيمة + 2، بسقف معقول)."""
    for col_idx in range(1, len(headers) + 1):
        longest = len(str(headers[col_idx - 1] or ""))
        for r in rows[:500]:
            if col_idx <= len(r):
                longest = max(longest, len(str(r[col_idx - 1] or "")))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(longest + 2, 60)


def export_xlsx(headers: List[str], rows: List[List[str]],
                kind: str = "tables") -> bytes:
    """Excel RTL منسق حسب العقد. kind=mechanic: ورقة موحدة + عمودا الفئة والمصدر."""
    headers = [str(h) for h in (headers or [])]
    rows = [list(r) for r in (rows or [])]

    if kind == "mechanic":
        for extra in MECHANIC_EXTRA_HEADERS:
            if extra not in headers:
                headers.append(extra)
        rows = [r + [""] * max(0, len(headers) - len(r)) for r in rows]

    wb = Workbook()
    ws = wb.active
    ws.title = "ميكانيك" if kind == "mechanic" else "جداول"
    ws.sheet_view.rightToLeft = True  # RTL

    if headers:
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
    for r in rows:
        ws.append(["" if c is None else str(c) for c in r])
    # حدود لكل الخلايا + محاذاة
    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 1),
                            max_col=max(len(headers), 1)):
        for cell in row:
            cell.border = BORDER
            if cell.row > 1:
                cell.alignment = Alignment(horizontal="right", vertical="center")
    if headers:
        ws.freeze_panes = "A2"  # تجميد صف الترويسة
        _auto_width(ws, headers, rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export(kind: str, fmt: str, headers: List[str], rows: List[List[str]]) -> tuple:
    """نقطة الدخول: يعيد (filename, bytes, media_type)."""
    kind = "mechanic" if kind == "mechanic" else "tables"
    ts = time.strftime("%Y%m%d_%H%M%S")
    if fmt == "xlsx":
        return (f"{kind}_{ts}.xlsx", export_xlsx(headers, rows, kind),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return (f"{kind}_{ts}.csv", export_csv(headers, rows), "text/csv; charset=utf-8")
